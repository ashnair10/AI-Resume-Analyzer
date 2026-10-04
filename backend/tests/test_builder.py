from io import BytesIO
import json
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from app.main import app
from src.resume_builder import ResumeExport, build_bundle, build_docx, filename_stem

CONTENT = """Example Candidate
AI Engineer | candidate@example.com | github.com/example
Summary
Build reliable document processing systems.
Technical Skills
Python, FastAPI, C++, RAG
Experience
AI Engineer | Example Company | 2022 - Present
- Built Python APIs and reduced processing time by 30%.
- Worked with operations teams to discover process steps and SLAs.
Projects
Resume analyzer | github.com/example/resume
- Developed FastAPI services for document parsing.
Education
B.Tech | Example University | 2022"""


def payload(**kwargs):
    return {"content": CONTENT, "company": "Example", "role": "AI Engineer", "reviewed": True, **kwargs}


def test_export_requires_review_and_valid_template():
    client = TestClient(app)
    assert client.post("/api/resume/export", json=payload(reviewed=False)).status_code == 400
    assert client.post("/api/resume/export", json=payload(template="unknown")).status_code == 422
    assert client.post("/api/resume/export?format=pdf", json=payload()).status_code == 400


@pytest.mark.parametrize("template", ["classic", "compact", "modern"])
def test_docx_download_has_no_tables_and_highlights_existing_terms(template):
    response = TestClient(app).post("/api/resume/export?format=docx", json=payload(template=template, highlight_terms=["Python", "C++", "invented-skill"]))
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    assert not document.tables
    text = "\n".join(p.text for p in document.paragraphs)
    assert "invented-skill" not in text
    assert "SLAs" in text
    assert any(run.bold and run.text == "Python" for p in document.paragraphs for run in p.runs)
    assert any(run.bold and run.text == "C++" for p in document.paragraphs for run in p.runs)
    assert response.headers["cache-control"] == "no-store"


def test_bundle_converts_docx_and_compares_extraction():
    bundle = build_bundle(ResumeExport(**payload(highlight_terms=["Python"])))
    with ZipFile(BytesIO(bundle)) as archive:
        manifest = json.loads(archive.read("checks.json"))
        assert manifest["extracted_text_matches"] is True
        assert manifest["pdf_pages"] == 1
        assert archive.read("resume-source.txt").decode() == CONTENT
        assert archive.read("Example_AI_Engineer.pdf").startswith(b"%PDF")
        assert "30%" in archive.read("pdf-extracted.txt").decode()


def test_pdf_failure_leaves_docx_download_working(monkeypatch):
    def failed(_):
        raise RuntimeError("converter missing")
    monkeypatch.setattr("src.resume_builder.convert_pdf", failed)
    client = TestClient(app)
    assert client.post("/api/resume/export", json=payload()).status_code == 503
    assert client.post("/api/resume/export?format=docx", json=payload()).status_code == 200


def test_import_and_review_work_without_ai():
    docx = build_docx(ResumeExport(**payload()))
    response = TestClient(app).post("/api/resume/import", files={"resume": ("resume.docx", docx)})
    assert response.status_code == 200
    assert "30%" in response.json()["content"]
    review = TestClient(app).post("/api/resume/review", json={"content": response.json()["content"], "job_description": "Seeking an AI Engineer with Python and FastAPI experience to build production APIs and reliable document processing systems.", "use_ai": False})
    assert review.status_code == 200
    assert review.json()["ai_review"] is None


def test_import_rejects_malformed_and_oversized_files():
    client = TestClient(app)
    assert client.post("/api/resume/import", files={"resume": ("broken.docx", b"not a docx")}).status_code == 422
    assert client.post("/api/resume/import", files={"resume": ("large.pdf", b"x" * (8 * 1024 * 1024 + 1))}).status_code == 413


def test_long_resume_preserves_all_content_and_warns():
    content = CONTENT + "\n" + "\n".join(f"- Built service {i} with Python and FastAPI for reliable document processing and human review workflows." for i in range(90))
    bundle = build_bundle(ResumeExport(**payload(content=content)))
    with ZipFile(BytesIO(bundle)) as archive:
        manifest = json.loads(archive.read("checks.json"))
        assert manifest["pdf_pages"] > 2
        assert manifest["warnings"]
        assert "service 89" in archive.read("pdf-extracted.txt").decode()


def test_file_names_cannot_escape_directories():
    assert "/" not in filename_stem("../../secret", "role\r\nheader")
    assert "\n" not in filename_stem("company", "role\r\nheader")


def test_ai_review_does_not_silently_truncate_input(monkeypatch):
    from core.ai_gateway import analyze_resume_ai
    monkeypatch.setenv("GEMINI_API_KEY", "test-key-not-used")
    result = analyze_resume_ai("x" * 20001, "Python role", {})
    assert not result.enabled
    assert "Shorten the input" in result.error
