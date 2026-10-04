"""One reviewed source -> DOCX -> PDF, with extraction verification.

No model rewrites are applied during export and no personal data is persisted.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, Field, field_validator
from .extractor import extract_resume

HEADINGS = {
    "summary", "professional summary", "profile", "objective", "experience",
    "professional experience", "work experience", "employment", "work history",
    "skills", "technical skills", "core competencies", "education", "projects",
    "personal projects", "selected projects", "certifications", "achievements",
    "awards", "publications", "languages", "volunteering", "interests",
}


class ResumeExport(BaseModel):
    content: str = Field(min_length=40, max_length=40000)
    company: str = Field(default="Application", max_length=120)
    role: str = Field(default="Resume", max_length=120)
    job_description: str = Field(default="", max_length=20000)
    template: Literal["classic", "compact"] = "compact"
    highlight_terms: list[str] = Field(default_factory=list, max_length=30)
    reviewed: bool = False

    @field_validator("content")
    @classmethod
    def clean_content(cls, value: str) -> str:
        if not value.strip() or any(ord(c) < 32 and c not in "\n\r\t" for c in value):
            raise ValueError("Resume text must be nonempty and contain no control characters.")
        return value.strip()

    @field_validator("highlight_terms")
    @classmethod
    def clean_terms(cls, values: list[str]) -> list[str]:
        if any(len(t) > 80 for t in values):
            raise ValueError("Each highlight term must be at most 80 characters.")
        return list(dict.fromkeys(t.strip() for t in values if t.strip()))


def filename_stem(company: str, role: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{company}_{role}").strip("_")
    return slug[:100] or "Application_Resume"


def add_emphasis(paragraph, text: str, terms: list[str]) -> None:
    if not terms:
        paragraph.add_run(text)
        return
    pattern = re.compile(r"(?<!\w)(?:" + "|".join(re.escape(t) for t in sorted(terms, key=len, reverse=True)) + r")(?!\w)", re.I)
    start = 0
    for match in pattern.finditer(text):
        paragraph.add_run(text[start:match.start()])
        paragraph.add_run(match.group()).bold = True
        start = match.end()
    paragraph.add_run(text[start:])


def build_docx(request: ResumeExport) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    margin = 0.65 if request.template == "compact" else 0.8
    section.top_margin = section.bottom_margin = Inches(margin)
    section.left_margin = section.right_margin = Inches(margin)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.5 if request.template == "compact" else 11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3 if request.template == "compact" else 5)
    normal.paragraph_format.line_spacing = 1.05 if request.template == "compact" else 1.12
    # Some default templates include inherited blue title borders. Strip them.
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    for name in ["Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(20 if name == "Title" else 11)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    lines = [line.strip() for line in request.content.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        heading = line.lstrip("# ").rstrip(":").strip()
        is_heading = heading.lower() in HEADINGS or line.startswith("## ")
        if index == 0:
            doc.add_paragraph(line, "Title")
        elif is_heading:
            doc.add_paragraph(heading, "Heading 1")
        else:
            bullet = re.match(r"^[-*•]\s+", line)
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            if bullet:
                p.paragraph_format.left_indent = Inches(.13)
                p.paragraph_format.first_line_indent = Inches(-.13)
            add_emphasis(p, ("- " + line[bullet.end():]) if bullet else line, request.highlight_terms)
    # Do not embed personal metadata or application notes in the resume.
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def convert_pdf(docx_bytes: bytes) -> bytes:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("PDF export needs LibreOffice. Use Docker Compose or install LibreOffice locally.")
    with tempfile.TemporaryDirectory(prefix="techcv-export-") as directory:
        root = Path(directory)
        source = root / "resume.docx"
        source.write_bytes(docx_bytes)
        # Each request has an isolated profile to avoid locks across simultaneous exports.
        subprocess.run(
            [executable, f"-env:UserInstallation={(root / 'profile').as_uri()}",
             "--headless", "--convert-to", "pdf", "--outdir", str(root), str(source)],
            check=True, timeout=60, capture_output=True,
        )
        target = root / "resume.pdf"
        if not target.exists():
            raise RuntimeError("LibreOffice did not produce a PDF. Check the server installation.")
        return target.read_bytes()


def normalized_text(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text)).replace("•", "-")


def build_bundle(request: ResumeExport) -> bytes:
    docx_bytes = build_docx(request)
    pdf_bytes = convert_pdf(docx_bytes)
    word = extract_resume(docx_bytes, "resume.docx")
    pdf = extract_resume(pdf_bytes, "resume.pdf")
    parity = normalized_text(word.text) == normalized_text(pdf.text)
    warnings = []
    if not parity:
        warnings.append("DOCX/PDF extracted text differs. Inspect both text files and the PDF before applying.")
    if pdf.page_count > 2:
        warnings.append("Resume is over two pages. Review relevance and length; content was not silently cut.")
    manifest = {
        "company": request.company, "role": request.role, "template": request.template,
        "source_sha256": hashlib.sha256(request.content.encode()).hexdigest(),
        "pdf_pages": pdf.page_count, "extracted_text_matches": parity,
        "warnings": warnings,
        "score_note": "Text checks do not predict any employer ATS score or selection outcome.",
    }
    stem = filename_stem(request.company, request.role)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr(f"{stem}.docx", docx_bytes)
        archive.writestr(f"{stem}.pdf", pdf_bytes)
        archive.writestr("checks.json", json.dumps(manifest, indent=2))
        archive.writestr("resume-source.txt", request.content)
        archive.writestr("docx-extracted.txt", word.text)
        archive.writestr("pdf-extracted.txt", pdf.text)
        archive.writestr("job-description.txt", request.job_description)
    return output.getvalue()
