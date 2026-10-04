from dataclasses import dataclass
from io import BytesIO
from collections import Counter
import fitz  # PyMuPDF
from docx import Document


@dataclass
class ResumeDocument:
    text: str
    page_count: int
    fonts: list[tuple[str, int]]
    font_sizes: list[tuple[float, int]]
    file_type: str


def extract_pdf(file_bytes: bytes) -> ResumeDocument:
    doc = fitz.open(stream=BytesIO(file_bytes), filetype="pdf")
    texts: list[str] = []
    font_counter: Counter[str] = Counter()
    size_counter: Counter[float] = Counter()

    for page in doc:
        texts.append(page.get_text("text"))
        data = page.get_text("dict")
        for block in data.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    font = str(span.get("font", "Unknown"))
                    size = round(float(span.get("size", 0)), 1)
                    font_counter[font] += len(span.get("text", "")) or 1
                    if size:
                        size_counter[size] += len(span.get("text", "")) or 1

    return ResumeDocument(
        text="\n".join(texts).strip(),
        page_count=len(doc),
        fonts=font_counter.most_common(8),
        font_sizes=size_counter.most_common(8),
        file_type="pdf",
    )


def extract_docx(file_bytes: bytes) -> ResumeDocument:
    document = Document(BytesIO(file_bytes))
    text_parts: list[str] = []
    font_counter: Counter[str] = Counter()
    size_counter: Counter[float] = Counter()

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
        for run in paragraph.runs:
            if run.font.name:
                font_counter[run.font.name] += len(run.text) or 1
            if run.font.size:
                size = round(run.font.size.pt, 1)
                size_counter[size] += len(run.text) or 1

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text.strip())

    return ResumeDocument(
        text="\n".join(text_parts).strip(),
        page_count=0,  # DOCX pagination depends on renderer; avoid pretending we know it.
        fonts=font_counter.most_common(8),
        font_sizes=size_counter.most_common(8),
        file_type="docx",
    )


def extract_resume(file_bytes: bytes, filename: str) -> ResumeDocument:
    if filename.lower().endswith(".pdf"):
        return extract_pdf(file_bytes)
    if filename.lower().endswith(".docx"):
        return extract_docx(file_bytes)
    raise ValueError("Unsupported file type. Upload a PDF or DOCX resume.")
